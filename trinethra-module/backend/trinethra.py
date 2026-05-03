import re
import json
import csv
import sys
from flask import Flask, request, jsonify
from flask_cors import CORS
import logging
import io
import base64
from docx import Document
from docx.shared import Inches

# Create a logger
logger = logging.getLogger(__name__)
logger.setLevel(logging.ERROR)

# Create a file handler
file_handler = logging.FileHandler('error.log')
file_handler.setLevel(logging.ERROR)

# Create a formatter and add it to the file handler
formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s')
file_handler.setFormatter(formatter)

# Add the file handler to the logger
logger.addHandler(file_handler)

# ...

# Modify the error handling code to log the error
def handle_exception(e):
    logger.error(f"An error occurred: {e}")
    return jsonify({'error': str(e)}), 500


app = Flask(__name__)
CORS(app)

try:
    from textblob import TextBlob
except ImportError:
    TextBlob = None

class TrinethraAssess:
    """Logic for DT Fellow performance evaluation."""

    POSITIVE_KEYWORDS = [
        'good', 'great', 'excellent', 'amazing', 'fantastic', 'wonderful', 'outstanding', 'superb',
        'reliable', 'trustworthy', 'helpful', 'valuable', 'useful', 'efficient', 'effective',
        'sincere', 'dedicated', 'committed', 'professional', 'proactive', 'initiative', 'hands-on',
        'happy', 'satisfied', 'pleased', 'glad', 'grateful', 'relieved', 'impressed',
        'best', 'strong', 'solid', 'improve', 'improved', 'saved', 'faster', 'better',
        'love', 'adore', 'appreciate', 'thank', 'thanks', 'perfect', 'flawless'
    ]

    NEGATIVE_KEYWORDS = [
        'bad', 'poor', 'terrible', 'awful', 'horrible', 'disappointing', 'inadequate', 'insufficient',
        'unreliable', 'untrustworthy', 'useless', 'ineffective', 'inefficient', 'careless',
        'lazy', 'slow', 'sloppy', 'inconsistent', 'unpredictable', 'unprofessional',
        'worried', 'concerned', 'upset', 'frustrated', 'annoyed', 'irritated', 'disappointed',
        'problem', 'issue', 'trouble', 'difficulty', 'struggle', 'fail', 'failed',
        'complaint', 'complain', 'reject', 'rejection', 'error', 'mistake', 'wrong',
        'worst', 'weak', 'lacking', 'missing', 'absent', 'late', 'delayed'
    ]

    NEGATORS = ['not', "n't", 'never', 'no', 'hardly', 'barely', 'rarely']

    KPI_KEYWORDS = {
        'lead_generation': ['new customers', 'new schools', 'leads', 'prospects'],
        'lead_conversion': ['closed', 'converted', 'signed'],
        'upselling': ['bigger', 'increased', 'selling more'],
        'cross_selling': ['additional', 'extra'],
        'nps': ['happy', 'satisfied', 'recommend'],
        'pat': ['waste', 'reduced', 'saved', 'profit'],
        'tat': ['faster', 'quicker', 'turnaround', 'speed'],
        'quality': ['rejection', 'defect', 'complaint', 'quality']
    }

    LAYER_1_KEYWORDS = ['helps', 'maintains', 'updates', 'handles', 'coordinates', 'assists', 'does', 'manages']
    LAYER_2_KEYWORDS = ['built', 'created', 'designed', 'started', 'set up', 'developed', 'automated', 'streamlined']

    DIMENSION_KEYWORDS = {
        'Driving Execution': ['task', 'delivers', 'on time', 'follow up', 'completes'],
        'Building Systems': ['system', 'tracker', 'sheet', 'template', 'process', 'automate'],
        'KPI Impact': ['faster', 'saved', 'reduced', 'increased', 'improved', 'metrics'],
        'Change Management': ['resistance', 'adopt', 'workers', 'floor team', 'compliance']
    }

    BIAS_INDICATORS = {
        'helpfulness': ['handles all', 'takes off my plate', 'big relief', 'helper', 'does everything', 'relieves', 'my right hand', "don't know how we managed"],
        'presence': ['always on', 'on the floor', 'physically', 'present', 'never leaves'],
        'halo': ['love', 'glowing', 'amazing', 'fantastic', 'couldn\'t manage without'],
        'recency': ['last week', 'recently', 'lately', 'these days', 'past few']
    }

    TASK_ABSORPTION_PHRASES = ['runs my', 'takes my calls', 'handles all my', 'takes over', 'doing raghav\'s']

    def assess_transcript(self, transcript):
        if not transcript or not isinstance(transcript, str):
            return {'error': 'Transcript is required', 'errors': ['Transcript is required']}

        text_lower = transcript.lower()
        evidence = self._extract_evidence(transcript)
        layers = self._detect_layers(text_lower)
        kpis = self._map_kpis(text_lower)
        dimensions = self._detect_dimensions(text_lower)
        biases = self._detect_biases(text_lower)

        score = self._calculate_score(text_lower, evidence, layers, dimensions)
        gaps = self._identify_gaps(dimensions, layers, evidence)
        questions = self._generate_questions(gaps, dimensions)

        return {
            'score': score,
            'evidence': evidence,
            'kpis': kpis,
            'dimensions': dimensions,
            'gaps': gaps,
            'questions': questions,
            'biases': biases,
            'layers': layers
        }

    def _extract_evidence(self, transcript):
        sentences = re.split(r'[.!?]+', transcript)
        evidence = []

        for sentence in sentences:
            sentence = sentence.strip()
            if len(sentence) < 20:
                continue

            polarity = self._sentence_polarity(sentence)
            sentiment = 'positive' if polarity > 0.1 else 'negative' if polarity < -0.1 else 'neutral'
            dimension = self._classify_dimension(sentence.lower())

            evidence.append({
                'quote': sentence[:200],
                'sentiment': sentiment,
                'dimension': dimension
            })

        return evidence[:6]

    def _classify_dimension(self, text):
        scores = {}
        for dim, keywords in self.DIMENSION_KEYWORDS.items():
            count = sum(1 for kw in keywords if kw in text)
            scores[dim] = count

        if not scores or max(scores.values()) == 0:
            return 'Driving Execution'
        return max(scores, key=scores.get)

    def _detect_layers(self, text):
        layer1_score = sum(1 for kw in self.LAYER_1_KEYWORDS if kw in text)
        layer2_score = sum(1 for kw in self.LAYER_2_KEYWORDS if kw in text)
        return {
            'execution': layer1_score > layer2_score,
            'systems_building': layer2_score >= layer1_score,
            'layer2_strength': layer2_score
        }

    def _map_kpis(self, text):
        matched_kpis = []
        for kpi, keywords in self.KPI_KEYWORDS.items():
            if any(kw in text for kw in keywords):
                matched_kpis.append(kpi.replace('_', ' ').title())
        return matched_kpis[:4]

    def _detect_dimensions(self, text):
        dimensions = {}
        for dim, keywords in self.DIMENSION_KEYWORDS.items():
            dimensions[dim] = any(kw in text for kw in keywords)
        return dimensions

    def _sentence_polarity(self, sentence):
        if TextBlob is not None:
            try:
                return TextBlob(sentence).sentiment.polarity
            except Exception:
                pass
        return self._keyword_polarity(sentence)

    def _keyword_polarity(self, text):
        """Fixed: Proper negation handling with correct return position"""
        text_lower = text.lower()
        words = re.findall(r'\b\w+\b', text_lower)
        score = 0
        negate = False
        
        for word in words:
            if word in self.NEGATORS:
                negate = True
                continue
            if word in self.POSITIVE_KEYWORDS:
                score += -1 if negate else 1
                negate = False
            elif word in self.NEGATIVE_KEYWORDS:
                score += 1 if negate else -1
                negate = False
            else:
                negate = False
        
        return score / max(len(words), 1)

    def _calculate_score(self, text, evidence, layers, dimensions):
        has_task_absorption = any(phrase in text for phrase in self.TASK_ABSORPTION_PHRASES)
        problem_identifier_phrases = ['noticed', 'found that', 'discovered', 'identified', 'quantified', 'tracked', 'analysis']
        has_problem_id = any(phrase in text for phrase in problem_identifier_phrases)
        has_system_creation = layers['systems_building'] and layers['layer2_strength'] >= 2
        lack_initiative = ['doesn\'t push back', 'doesn\'t question', 'just does', 'no initiative', 'does what i tell', 'tells him to do']
        has_lack_initiative = any(phrase in text for phrase in lack_initiative)
        dimension_count = sum(dimensions.values())

        if has_lack_initiative:
            core_score = 5
        elif has_problem_id and has_system_creation:
            core_score = 8
        elif has_problem_id:
            core_score = 7
        elif has_system_creation:
            core_score = 7 if dimension_count >= 2 else 6
        elif dimension_count >= 2:
            core_score = 6
        else:
            core_score = 5

        if has_task_absorption and core_score > 6:
            core_score = 6

        if core_score <= 3:
            band = 'Need Attention'
            label = 'Motivated but Directionless' if core_score == 3 else 'Consistent Performer'
        elif core_score <= 6:
            band = 'Productivity'
            labels = {4: 'Careless and Inconsistent', 5: 'Consistent Performer', 6: 'Reliable and Productive'}
            label = labels.get(core_score, 'Consistent Performer')
        else:
            band = 'Performance'
            labels = {7: 'Problem Identifier', 8: 'Problem Solver', 9: 'Innovative and Experimental', 10: 'Exceptional Performer'}
            label = labels.get(core_score, 'Problem Identifier')

        justification = self._generate_justification(core_score, text, layers, dimensions)

        return {
            'value': core_score,
            'label': label,
            'band': band,
            'justification': justification
        }

    def _generate_justification(self, score, text, layers, dimensions):
        justifications = {
            5: "Shows reliable task execution but limited systems building.",
            6: "Consistent performer with some operational responsibilities.",
            7: "Shows problem identification with evidence of tracking or analysis.",
            8: "Demonstrates systems building with tools or processes created."
        }
        base = justifications.get(score, "Mixed evidence in transcript.")
        
        if not layers['systems_building']:
            base += " No clear systems building detected."
        if not dimensions.get('Change Management', False):
            base += " No change management evidence present."
        return base

    def _identify_gaps(self, dimensions, layers, evidence):
        """Fixed: Uses original comprehensive gap detection"""
        gaps = []
        if not layers['systems_building']:
            gaps.append({'dimension': 'systems_building', 'detail': 'No evidence of systems, processes, or tools created'})
        if not dimensions.get('Change Management', False):
            gaps.append({'dimension': 'change_management', 'detail': 'No mention of how Fellow handles resistance or gets team adoption'})
        if not dimensions.get('KPI Impact', False):
            gaps.append({'dimension': 'kpi_impact', 'detail': 'No measurable business outcome connections'})
        if not dimensions.get('Building Systems', False):
            gaps.append({'dimension': 'building_systems', 'detail': 'No structures, trackers, or SOPs documented'})
        return gaps

    def _generate_questions(self, gaps, dimensions):
        questions = []
        for gap in gaps:
            dim = gap['dimension']
            if dim == 'systems_building':
                questions.append('If the Fellow took a week off, what would stop working? What would keep running on its own?')
            elif dim == 'change_management':
                questions.append('How do the floor workers respond when the Fellow asks them to do something differently?')
            elif dim == 'kpi_impact':
                questions.append('Since the Fellow started, has any specific number improved — speed, complaints, rejections?')
            elif dim == 'building_systems':
                questions.append('Has the Fellow built anything that your team uses regularly after they leave?')

        if len(questions) < 3:
            questions.append('When was the last time the Fellow suggested a new approach or process?')
        return questions[:3]

    def _detect_biases(self, text):
        detected_biases = []
        for bias, indicators in self.BIAS_INDICATORS.items():
            if any(ind in text for ind in indicators):
                detected_biases.append(bias)
        return detected_biases

    def analyze_sentiment_keywords(self, text):
        """Detailed sentiment keyword analysis"""
        text_lower = text.lower()
        positive_count = sum(1 for kw in self.POSITIVE_KEYWORDS if kw in text_lower)
        negative_count = sum(1 for kw in self.NEGATIVE_KEYWORDS if kw in text_lower)
        keywords_found = []

        for kw in self.POSITIVE_KEYWORDS:
            if kw in text_lower:
                keywords_found.append({'keyword': kw, 'type': 'positive'})
        for kw in self.NEGATIVE_KEYWORDS:
            if kw in text_lower:
                keywords_found.append({'keyword': kw, 'type': 'negative'})

        has_negation = any(neg in text_lower for neg in self.NEGATORS)
        total = positive_count + negative_count

        if total == 0:
            return {'score': 0.5, 'label': 'neutral', 'keywords_found': [], 'confidence': 'low'}

        if positive_count > negative_count:
            raw_score = (positive_count / total) * 0.5 + 0.5
        elif negative_count > positive_count:
            raw_score = 0.5 - (negative_count / total) * 0.5
        else:
            raw_score = 0.5

        if has_negation and total > 0:
            raw_score = 1.0 - raw_score

        score = max(0.0, min(1.0, raw_score))
        label = 'positive' if score >= 0.6 else 'negative' if score <= 0.4 else 'neutral'
        confidence = 'high' if total >= 5 else 'medium' if total >= 2 else 'low'

        return {
            'score': round(score, 2),
            'label': label,
            'keywords_found': keywords_found[:10],
            'confidence': confidence,
            'positive_count': positive_count,
            'negative_count': negative_count
        }


class TrinethraCore:
    """Handles text cleaning, batch processing, and API integration."""
    
    SPECIAL_CHARACTERS = r'[^a-zA-Z0-9\s.,!?\'"-]'

    def __init__(self):
        self.processor = TrinethraAssess()

    def process_json_batch(self, json_data):
        all_results = []
        transcripts_list = json_data.get("transcripts", [])
        
        for entry in transcripts_list:
            raw_text = entry.get("transcript", "")
            if len(raw_text) < 5:
                continue
                
            fellow_info = entry.get("fellow", {})
            name = fellow_info.get("name", "Unknown")
            
            analysis = self.processor.assess_transcript(raw_text)
            analysis['fellow_name'] = name
            analysis['id'] = entry.get("id")
            analysis['transcript'] = raw_text  # Add original transcript for summary
            all_results.append(analysis)
            
        return all_results

    def preprocess_text(self, text):
        if not text or not isinstance(text, str):
            return ""
        cleaned = re.sub(self.SPECIAL_CHARACTERS, '', text)
        standardized = cleaned.lower()
        normalized = re.sub(r'\s+', ' ', standardized).strip()
        return normalized

    def process_feedback(self, feedback):
        if not feedback or not isinstance(feedback, str):
            return {"error": "Invalid feedback provided"}
        clean_text = self.preprocess_text(feedback)
        analysis = self.processor.assess_transcript(clean_text)
        analysis['transcript'] = feedback  # Add original transcript for summary
        return analysis


# Global instance
core_instance = TrinethraCore()

@app.route('/api/analyze', methods=['POST'])
def analyze():
    """Batch analysis endpoint"""
    data = request.get_json(force=True, silent=True)
    
    if not data or 'transcripts' not in data:
        return jsonify({"error": "Payload must contain 'transcripts' array"}), 400

    try:
        results = core_instance.process_json_batch(data)
        return jsonify(results)
    except Exception as e:
        print(f"ERROR: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/analyze_single', methods=['POST'])
def analyze_single():
    """Single transcript analysis endpoint"""
    data = request.get_json(force=True, silent=True)
    
    if not data or 'transcript' not in data:
        return jsonify({"error": "Payload must contain 'transcript' field"}), 400

    try:
        result = core_instance.process_feedback(data['transcript'])
        return jsonify(result)
    except Exception as e:
        print(f"ERROR: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/save_review', methods=['POST'])
def save_review():
    """Save user review to CSV and JSON"""
    data = request.get_json()
    if not data:
        return jsonify({'error': 'No data provided'}), 400
    
    review = data.get('review', {})
    timestamp = data.get('timestamp', '')
    
    # Append to CSV
    with open('test_feedback.csv', 'a', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=['id', 'status', 'score', 'reviewer', 'reject_reason', 'kpis', 'gaps'])
        if f.tell() == 0:  # Write header if empty
            writer.writeheader()
        writer.writerow({
            'id': review.get('id'),
            'status': review.get('status', 'draft'),
            'score': review.get('edits', {}).get('score.value'),
            'reviewer': review.get('reviewer', ''),
            'reject_reason': review.get('rejectReason', ''),
            'kpis': json.dumps(review.get('data', {}).get('kpis', [])),
            'gaps': json.dumps(review.get('data', {}).get('gaps', []))
        })
    
    # Append to reviews.json
    reviews = []
    try:
        with open('reviews.json', 'r') as f:
            reviews = json.load(f)
    except FileNotFoundError:
        pass
    reviews.append({**review, 'timestamp': timestamp})
    with open('reviews.json', 'w') as f:
        json.dump(reviews, f, indent=2)
    
    return jsonify({'success': True, 'message': 'Review saved to CSV and JSON'})

@app.route('/api/reviews', methods=['GET'])
def get_reviews():
    """Load saved reviews"""
    try:
        with open('reviews.json', 'r') as f:
            return jsonify(json.load(f))
    except FileNotFoundError:
        return jsonify([])

@app.route('/api/export_word', methods=['POST'])
def export_word():
    """Generate Word document for current analysis and reviews"""
    data = request.get_json()
    if not data:
        return jsonify({'error': 'No data provided'}), 400
    
    analysis = data.get('analysis', [])
    reviews = data.get('reviews', [])
    reviewer = data.get('reviewerName', 'Anonymous')
    
    if not analysis:
        return jsonify({'error': 'No analysis data'}), 400
    
    # Create DOCX
    doc = Document()
    doc.add_heading('Trinethra Performance Report', 0)
    doc.add_paragraph(f'Reviewer: {reviewer}')
    doc.add_paragraph(f'Generated: {data.get("timestamp", "Now")}')
    doc.add_page_break()
    
    # Single or multiple
    if len(analysis) == 1:
        item = analysis[0]
        review_item = reviews[0] if reviews else {}
        _add_fellow_report(doc, item, review_item)
    else:
        for i, item in enumerate(analysis):
            review_item = reviews[i] if i < len(reviews) else {}
            doc.add_heading(f'Fellow {i+1}: {item.get("fellow_name", "Unnamed")}', level=1)
            _add_fellow_report(doc, item, review_item)
            doc.add_page_break()
    
    # Save to bytes
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    docx_base64 = base64.b64encode(stream.read()).decode('utf-8')
    
    return jsonify({'docx_base64': docx_base64})


def _add_fellow_report(doc, analysis, review):
    """Helper to add single fellow report"""
    edits = review.get('edits', {})
    
    # Score
    doc.add_heading('Performance Score', level=2)
    score_val = edits.get('score.value', analysis.get('score', {}).get('value', 'N/A'))
    score_label = edits.get('score.label', analysis.get('score', {}).get('label', 'N/A'))
    p = doc.add_paragraph()
    p.add_run(f'Score: {score_val}/10 (').bold = True
    p.add_run(f'{score_label})').bold = False
    doc.add_paragraph(edits.get('score.justification', analysis.get('score', {}).get('justification', '')))
    
    # KPIs
    if analysis.get('kpis'):
        doc.add_heading('KPIs Impacted', level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'KPI'
        hdr_cells[1].text = 'Status'
        for kpi in analysis['kpis']:
            row_cells = table.add_row().cells
            row_cells[0].text = kpi
            row_cells[1].text = 'Active'  # Simple
        doc.add_paragraph('(KPIs detected from transcript analysis)')
    
    # Gaps
    if analysis.get('gaps'):
        doc.add_heading('Identified Gaps', level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Dimension'
        hdr_cells[1].text = 'Detail'
        for gap in analysis['gaps']:
            row_cells = table.add_row().cells
            row_cells[0].text = gap.get('dimension', '')
            row_cells[1].text = gap.get('detail', '')
    
    # Questions
    if analysis.get('questions'):
        doc.add_heading('Follow-up Questions', level=2)
        for q in analysis['questions']:
            doc.add_paragraph(q, style='List Bullet')
    
    # Evidence
    if analysis.get('evidence'):
        doc.add_heading('Key Evidence', level=2)
        for ev in analysis['evidence']:
            p = doc.add_paragraph(f'"{ev["quote"]} ({ev["sentiment"]} - {ev["dimension"]})')
    
    # Status from review
    status = review.get('status', 'draft')
    doc.add_paragraph(f'\nReview Status: {status.upper()}', style='Intense Quote')


@app.route('/api/generate_summary', methods=['POST'])
def generate_summary_endpoint():
    """Generate LLM summary for transcript - with robust fallback"""
    data = request.get_json(force=True, silent=True)
    print(f"SUMMARY REQUEST: transcript length={len(data.get('transcript', '')) if data else 0}")
    
    if not data or 'transcript' not in data:
        print("SUMMARY ERROR: Missing transcript")
        return jsonify({"error": "Payload must contain 'transcript' field"}), 400
    
    transcript = (data['transcript'] or '').strip()
    if not transcript:
        print("SUMMARY ERROR: Empty transcript")
        return jsonify({'summary': 'No content - summary unavailable.'}), 400
    
    try:
        from summary_chain import generate_summary
        result = generate_summary(transcript)
        print(f"SUMMARY SUCCESS: {result[:50]}...")
        return jsonify({'summary': result})
    except ImportError as e:
        print(f"SUMMARY FALLBACK (import): {e}")
    except Exception as e:
        print(f"SUMMARY FALLBACK (error): {str(e)}")
    
    # Robust fallback: first sentences or truncated
    sentences = [s.strip() for s in transcript.split('.') if s.strip()]
    fallback = '. '.join(sentences[:3]) + '.' if sentences else transcript[:300]
    summary = fallback[:300] + '...' if len(fallback) > 300 else fallback
    print(f"SUMMARY FALLBACK: {summary[:50]}...")
    return jsonify({'summary': summary})


@app.route('/health', methods=['GET'])
def health():
    """Health check endpoint"""
    return jsonify({"status": "healthy", "service": "Trinethra"})


if __name__ == '__main__':
    print("Trinethra Module: Ready on Port 5000")
    print("Endpoints:")
    print("  POST /api/analyze - Batch analysis")
    print("  POST /api/analyze_single - Single transcript")
    print("  POST /api/save_review - Save user review")
    print("  GET /api/reviews - Load reviews")
    print("  POST /api/generate_summary - Generate transcript summary")
    print("  GET /health - Health check")
    app.run(host='0.0.0.0', port=5000, debug=True)
